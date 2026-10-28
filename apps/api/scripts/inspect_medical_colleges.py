import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\Likhith Kumar\Music\APPLYGIELDS\India_Top_50_Medical_Colleges.xlsx', data_only=True)
ws = wb['Medical Colleges']
headers = [str(ws.cell(1, col).value) for col in range(1, ws.max_column + 1)]

print(f"Total rows: {ws.max_row}")
none_counts = {'loc': 0, 'type': 0, 'estd': 0, 'about': 0, 'courses': 0, 'admissions': 0, 'placements': 0, 'cutoffs': 0}

for r in range(2, ws.max_row + 1):
    pos = ws.cell(r, 1).value
    name = str(ws.cell(r, 2).value or '').encode('ascii', 'replace').decode('ascii')
    fee = str(ws.cell(r, 3).value or '').encode('ascii', 'replace').decode('ascii')
    sal = str(ws.cell(r, 4).value or '').encode('ascii', 'replace').decode('ascii')
    loc = ws.cell(r, 9).value
    inst_type = ws.cell(r, 10).value
    estd = ws.cell(r, 11).value
    about = ws.cell(r, 12).value
    courses = ws.cell(r, 13).value
    admissions = ws.cell(r, 14).value
    placements = ws.cell(r, 15).value
    cutoffs = ws.cell(r, 16).value
    
    if not loc: none_counts['loc'] += 1
    if not inst_type: none_counts['type'] += 1
    if not estd: none_counts['estd'] += 1
    if not about: none_counts['about'] += 1
    if not courses: none_counts['courses'] += 1
    if not admissions: none_counts['admissions'] += 1
    if not placements: none_counts['placements'] += 1
    if not cutoffs: none_counts['cutoffs'] += 1
    
    has_loc = "OK" if loc else "NONE"
    has_type = "OK" if inst_type else "NONE"
    has_estd = "OK" if estd else "NONE"
    has_about = "OK" if about else "NONE"
    
    print(f"Row {r:2d} | #{pos:2} | {name[:35]:35} | {fee[:15]:15} | Loc:{has_loc:4} | Type:{has_type:4} | Estd:{has_estd:4} | About:{has_about:4}")

print("\nNone counts:", none_counts)
